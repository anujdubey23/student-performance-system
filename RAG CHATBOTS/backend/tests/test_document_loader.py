import pytest
import os
import tempfile
import docx
import fitz
from backend.app.rag.document_loader import DocumentLoader

def test_load_txt():
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as f:
        f.write("Hello RAGify! This is a plain text test document for retrieval.")
        temp_path = f.name

    try:
        pages = DocumentLoader.load(temp_path, "sample.txt")
        assert len(pages) == 1
        assert "Hello RAGify!" in pages[0]["text"]
        assert pages[0]["page_number"] == 1
        assert pages[0]["source"] == "sample.txt"
    finally:
        os.remove(temp_path)

def test_load_docx():
    doc = docx.Document()
    doc.add_heading("Deep Learning Overview", level=1)
    doc.add_paragraph("Neural networks learn hierarchical feature representations through backpropagation.")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_path = f.name
    doc.save(temp_path)

    try:
        pages = DocumentLoader.load(temp_path, "deep_learning.docx")
        assert len(pages) == 1
        assert "Neural networks learn hierarchical" in pages[0]["text"]
        assert pages[0]["document_name"] == "deep_learning.docx"
    finally:
        os.remove(temp_path)

def test_load_pdf_with_pages():
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 72), "Page 1: Transformers revolutionized Natural Language Processing.")
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Page 2: Self-attention enables parallelized sequence training.")

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        temp_path = f.name
    doc.save(temp_path)
    doc.close()

    try:
        pages = DocumentLoader.load(temp_path, "transformers.pdf")
        assert len(pages) == 2
        assert pages[0]["page_number"] == 1
        assert "Transformers revolutionized" in pages[0]["text"]
        assert pages[1]["page_number"] == 2
        assert "Self-attention enables" in pages[1]["text"]
    finally:
        os.remove(temp_path)

def test_empty_file_handling():
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w") as f:
        f.write("")
        temp_path = f.name

    try:
        with pytest.raises(ValueError):
            DocumentLoader.load(temp_path, "empty.txt")
    finally:
        os.remove(temp_path)
