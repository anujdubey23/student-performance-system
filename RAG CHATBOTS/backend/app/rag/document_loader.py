import os
from typing import List, Dict, Any
from backend.app.utils.logging import logger

class DocumentLoader:
    """
    Modular document loader supporting PDF, DOCX, and TXT files.
    Preserves page numbers (especially for PDF) and document metadata.
    """

    @classmethod
    def load(cls, file_path: str, filename: str) -> List[Dict[str, Any]]:
        """
        Loads document content by file extension.
        Returns a list of page/section dictionaries:
        [
            {
                "text": str,
                "page_number": int or None,
                "source": str,
                "document_name": str
            }
        ]
        """
        ext = os.path.splitext(filename)[1].lower()

        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found at path: {file_path}")

        if ext == ".pdf":
            return cls._load_pdf(file_path, filename)
        elif ext == ".docx":
            return cls._load_docx(file_path, filename)
        elif ext == ".txt":
            return cls._load_txt(file_path, filename)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    @classmethod
    def _load_pdf(cls, file_path: str, filename: str) -> List[Dict[str, Any]]:
        """Extracts text per page from PDF using PyMuPDF (fitz)."""
        import fitz  # PyMuPDF

        pages = []
        try:
            doc = fitz.open(file_path)
            if doc.is_encrypted:
                raise ValueError(f"PDF '{filename}' is password protected / encrypted.")

            if len(doc) == 0:
                raise ValueError(f"PDF '{filename}' has no pages.")

            for page_idx in range(len(doc)):
                page = doc[page_idx]
                page_text = page.get_text("text") or ""
                # Page numbers in UI should be 1-indexed
                pages.append({
                    "text": page_text,
                    "page_number": page_idx + 1,
                    "source": filename,
                    "document_name": filename
                })

            doc.close()
            logger.info(f"Loaded {len(pages)} pages from PDF '{filename}'")
            return pages
        except Exception as e:
            logger.error(f"Failed to load PDF '{filename}': {str(e)}")
            raise

    @classmethod
    def _load_docx(cls, file_path: str, filename: str) -> List[Dict[str, Any]]:
        """Extracts text from DOCX using python-docx."""
        import docx

        try:
            doc = docx.Document(file_path)
            full_text = []

            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())

            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))

            combined_text = "\n\n".join(full_text)
            if not combined_text.strip():
                raise ValueError(f"DOCX '{filename}' contains no readable text.")

            logger.info(f"Loaded DOCX '{filename}' ({len(combined_text)} characters)")
            return [{
                "text": combined_text,
                "page_number": 1,
                "source": filename,
                "document_name": filename
            }]
        except Exception as e:
            logger.error(f"Failed to load DOCX '{filename}': {str(e)}")
            raise

    @classmethod
    def _load_txt(cls, file_path: str, filename: str) -> List[Dict[str, Any]]:
        """Reads plain text file with UTF-8 or fallback encodings."""
        encodings = ["utf-8", "latin-1", "cp1252"]
        content = None

        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if content is None:
            raise ValueError(f"Could not decode TXT file '{filename}' with supported encodings.")

        if not content.strip():
            raise ValueError(f"TXT file '{filename}' is empty.")

        logger.info(f"Loaded TXT '{filename}' ({len(content)} characters)")
        return [{
            "text": content,
            "page_number": 1,
            "source": filename,
            "document_name": filename
        }]
